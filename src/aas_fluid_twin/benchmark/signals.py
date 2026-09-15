"""The plant's signal dictionary.

``Simulation_Variable_Mapping.xlsx`` in the benchmark is the authoritative record of
every signal: its tag, its OPC UA node, its unit, and the Modelica variable it
corresponds to. This module parses it into typed objects, applies a small set of
*documented* corrections, and enriches it with instrument type data from the vendor
datasheets (``resources/instruments.yaml``).

Two rules govern the parsing:

* **The CSV column name is a machine key and is never altered.** It joins the AAS, the
  CSV files and the time-series store, so the upstream typo ``Tempreature_before_Pump_P201``
  survives verbatim in :attr:`Signal.channel`. Corrected spelling lives in
  :attr:`Signal.display_name`.
* **Every correction is declared, not silent.** See :data:`CORRECTIONS`; each applied
  correction is recorded on the dictionary and asserted by the tests.
"""

from __future__ import annotations

import enum
from collections.abc import Iterator, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml

from aas_fluid_twin import config

__all__ = [
    "CORRECTIONS",
    "ChannelGroup",
    "Correction",
    "DeviceType",
    "Quality",
    "Role",
    "Signal",
    "SignalDictionary",
    "SignalKind",
    "load_signal_dictionary",
]


class Role(enum.StrEnum):
    SENSOR = "Sensor"
    ACTUATOR = "Actuator"


class SignalKind(enum.StrEnum):
    BINARY = "binary"
    ANALOGUE = "analogue"


class DeviceType(enum.StrEnum):
    LEVEL = "level"
    PRESSURE = "pressure"
    TEMPERATURE = "temperature"
    FLOW = "flow"
    VALVE = "valve"
    PUMP = "pump"
    MIXER = "mixer"


class Quality(enum.StrEnum):
    GOOD = "good"
    UNRELIABLE = "unreliable"


class ChannelGroup(enum.StrEnum):
    """Channels that share a unit and a physical meaning, i.e. one plot and one y-axis."""

    LEVEL_SWITCH = "level_switch"
    VOLUME = "volume"
    PRESSURE = "pressure"
    TEMPERATURE = "temperature"
    FLOW = "flow"
    VALVE = "valve"
    DRIVE = "drive"
    LEVEL_ULTRASONIC = "level_ultrasonic"
    LEVEL_FROM_VOLUME = "level_from_volume"
    LEVEL_FROM_PRESSURE = "level_from_pressure"
    LEVEL_FROM_PRESSURE_REL = "level_from_pressure_rel"


#: Values the mapping table uses to mean "nothing here".
_SENTINELS = frozenset({"", "binary", "not_assigned", "not_needed", "R201_not_in_model", "None"})

_EXPECTED_HEADER: tuple[str, ...] = (
    "ID",
    "Sensor_ID",
    "Role",
    "Binary_or_Analogue",
    "Type_of_Sensor_or_Actuator",
    "Name",
    "Sensor_OPCUA_Node_ID",
    "Unit_Sensor",
    "Simulation_variable_name",
    "Unit_Simulation_variable",
    "Simulation_initialization_parameters",
    "Unit_Simulation_initialization_parameters",
)


@dataclass(frozen=True, slots=True)
class Correction:
    """A declared, reviewable change to what the mapping table says."""

    key: str
    field_name: str
    was: str
    now: str
    reason: str


#: Corrections applied to the upstream mapping table. Each is justified in
#: ``docs/benchmark-deviations.md``; the tests pin this list so none can creep in unnoticed.
CORRECTIONS: tuple[Correction, ...] = (
    Correction(
        key="Tank_B204_level_calculated_via_VolumeB204",
        field_name="sensor_id",
        was="Level_B203",
        now="Level_B204",
        reason=(
            "Mapping table row 39 repeats the B203 sensor id while describing B204 "
            "(Simulation_variable_name = tank_B204.level). See deviation N2."
        ),
    ),
)

#: Human-readable names where the CSV column is misspelled or terse. The column itself is
#: never renamed — see the module docstring.
_DISPLAY_NAMES: Mapping[str, str] = {
    "Tempreature_before_Pump_P201": "Temperature before pump P201",
    "Temperature_in_B204": "Temperature in B204",
}

#: CSV-column suffix or prefix -> channel group. Order matters: the first match wins, so the
#: more specific ``_until_zero_level`` variant must be tested before the plain PI variant.
_GROUP_RULES: tuple[tuple[str, ChannelGroup], ...] = (
    ("_until_zero_level", ChannelGroup.LEVEL_FROM_PRESSURE_REL),
    ("_level_calculated_via_LI", ChannelGroup.LEVEL_ULTRASONIC),
    ("_level_calculated_via_Volume", ChannelGroup.LEVEL_FROM_VOLUME),
    ("_level_calculated_via_PI", ChannelGroup.LEVEL_FROM_PRESSURE),
    ("_Volume", ChannelGroup.VOLUME),
    ("Pressure_below_", ChannelGroup.PRESSURE),
    ("Flow_after_", ChannelGroup.FLOW),
    ("Valve_", ChannelGroup.VALVE),
)


@dataclass(frozen=True, slots=True)
class Signal:
    """One row of the plant's signal dictionary."""

    index: int
    sensor_id: str
    role: Role
    kind: SignalKind
    device_type: DeviceType
    channel: str
    """CSV column name, verbatim. The join key across AAS, CSV and the time-series store."""
    display_name: str
    opcua_node_id: str
    unit: str | None
    sim_variable: str | None
    sim_unit: str | None
    sim_init_parameter: str | None
    group: ChannelGroup
    asset_tag: str | None
    quality: Quality = Quality.GOOD
    quality_reason: str | None = None
    instrument: Mapping[str, Any] | None = None
    """Type data from the vendor datasheet, or ``None`` where no datasheet covers it."""

    @property
    def is_binary(self) -> bool:
        return self.kind is SignalKind.BINARY

    @property
    def is_mapped_to_simulation(self) -> bool:
        return self.sim_variable is not None

    @property
    def measuring_span(self) -> Mapping[str, Any] | None:
        if self.instrument is None:
            return None
        span = self.instrument.get("measuring_span")
        return span if isinstance(span, dict) else None


@dataclass(frozen=True, slots=True)
class SignalDictionary(Sequence[Signal]):
    """All plant signals, with lookups by the keys that actually get used."""

    signals: tuple[Signal, ...]
    corrections_applied: tuple[Correction, ...] = ()
    _by_channel: dict[str, Signal] = field(default_factory=dict, repr=False)
    _by_sensor_id: dict[str, Signal] = field(default_factory=dict, repr=False)

    def __post_init__(self) -> None:
        for sig in self.signals:
            if sig.channel in self._by_channel:
                raise ValueError(f"duplicate CSV column in the mapping table: {sig.channel!r}")
            self._by_channel[sig.channel] = sig
            self._by_sensor_id.setdefault(sig.sensor_id, sig)

    # Sequence protocol -------------------------------------------------------
    def __len__(self) -> int:
        return len(self.signals)

    def __getitem__(self, index: int) -> Signal:  # type: ignore[override]
        return self.signals[index]

    def __iter__(self) -> Iterator[Signal]:
        return iter(self.signals)

    # Lookups -----------------------------------------------------------------
    def by_channel(self, channel: str) -> Signal:
        try:
            return self._by_channel[channel]
        except KeyError:
            raise KeyError(f"no signal for CSV column {channel!r}") from None

    def by_sensor_id(self, sensor_id: str) -> Signal:
        try:
            return self._by_sensor_id[sensor_id]
        except KeyError:
            raise KeyError(f"no signal with Sensor_ID {sensor_id!r}") from None

    def in_group(self, group: ChannelGroup) -> tuple[Signal, ...]:
        return tuple(s for s in self.signals if s.group is group)

    def for_asset(self, tag: str) -> tuple[Signal, ...]:
        return tuple(s for s in self.signals if s.asset_tag == tag)

    def with_role(self, role: Role) -> tuple[Signal, ...]:
        return tuple(s for s in self.signals if s.role is role)

    @property
    def channels(self) -> tuple[str, ...]:
        return tuple(s.channel for s in self.signals)

    @property
    def actuator_channels(self) -> tuple[str, ...]:
        return tuple(s.channel for s in self.signals if s.role is Role.ACTUATOR)


# --- parsing -----------------------------------------------------------------


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _nullable(value: object) -> str | None:
    text = _clean(value)
    if not text or text in _SENTINELS or text.startswith("not_needed"):
        return None
    return text


def _classify_group(channel: str, device_type: DeviceType) -> ChannelGroup:
    for needle, group in _GROUP_RULES:
        if needle in channel:
            return group
    if device_type is DeviceType.TEMPERATURE:
        return ChannelGroup.TEMPERATURE
    if device_type in (DeviceType.PUMP, DeviceType.MIXER):
        return ChannelGroup.DRIVE
    if device_type is DeviceType.LEVEL:
        return ChannelGroup.LEVEL_SWITCH
    raise ValueError(f"cannot classify channel {channel!r} of type {device_type}")


def _asset_tag(channel: str, sensor_id: str) -> str | None:
    """The plant tag the signal belongs to (``B201``, ``P201``, ``V204`` …)."""
    for tag in ("B201", "B202", "B203", "B204", "P201", "P202", "R201"):
        if tag in channel:
            return tag
    if sensor_id.startswith("V") and sensor_id[1:].isdigit():
        return sensor_id
    return None


@dataclass(frozen=True, slots=True)
class _DerivedQuality:
    """Quality of a channel the PLC computes, inherited from its source instruments."""

    quality: Quality
    derived_from: tuple[str, ...]
    reason: str | None


def _load_instruments(
    path: Path | None = None,
) -> tuple[dict[str, dict[str, Any]], dict[ChannelGroup, _DerivedQuality]]:
    """Read ``resources/instruments.yaml`` into per-sensor and per-group lookups."""
    source = path or (config.RESOURCE_DIR / "instruments.yaml")
    raw: dict[str, Any] = yaml.safe_load(source.read_text(encoding="utf-8"))

    by_sensor: dict[str, dict[str, Any]] = {}
    for family, record in raw.get("families", {}).items():
        entry = {k: v for k, v in record.items() if k != "applies_to"}
        entry["family"] = family
        for sensor_id in record.get("applies_to", ()):
            by_sensor[sensor_id] = entry

    by_group: dict[ChannelGroup, _DerivedQuality] = {}
    for group_key, record in (raw.get("derived_quality") or {}).items():
        by_group[ChannelGroup(group_key)] = _DerivedQuality(
            quality=Quality(str(record["quality"])),
            derived_from=tuple(str(s) for s in record.get("derived_from", ())),
            reason=(str(record["reason"]).strip() if record.get("reason") else None),
        )
    return by_sensor, by_group


def _resolve_quality(
    sensor_id: str,
    group: ChannelGroup,
    instrument: Mapping[str, Any] | None,
    derived: Mapping[ChannelGroup, _DerivedQuality],
    instruments: Mapping[str, Mapping[str, Any]],
) -> tuple[Quality, str | None]:
    """Quality of a channel, and why.

    A channel measured directly takes its instrument's quality. A channel the PLC *derives*
    takes the quality of the instruments it is computed from — which is the case that matters
    here, because the pressure-derived levels carry no instrument of their own yet inherit the
    pressure transmitters' static-pressure limitation.
    """
    rule = derived.get(group)
    if rule is not None:
        reason = rule.reason
        if reason is None and rule.derived_from:
            source = instruments.get(rule.derived_from[0], {})
            reason = source.get("quality_reason")
        return rule.quality, reason if rule.quality is not Quality.GOOD else None

    if instrument is not None:
        quality = Quality(str(instrument.get("quality", "good")))
        return quality, instrument.get("quality_reason") if quality is not Quality.GOOD else None

    return Quality.GOOD, None


def _read_rows(path: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = tuple(_clean(cell) for cell in next(rows))
        if header[: len(_EXPECTED_HEADER)] != _EXPECTED_HEADER:
            raise ValueError(
                "unexpected mapping-table header; the benchmark file changed shape.\n"
                f"  expected: {_EXPECTED_HEADER}\n  got:      {header}"
            )
        return [
            {key: _clean(cell) for key, cell in zip(header, row, strict=False)}
            for row in rows
            if any(cell is not None for cell in row)
        ]
    finally:
        workbook.close()


def load_signal_dictionary(
    mapping_file: Path | None = None,
    *,
    instruments_file: Path | None = None,
) -> SignalDictionary:
    """Parse the benchmark's variable mapping table into a :class:`SignalDictionary`."""
    path = mapping_file or config.VARIABLE_MAPPING_FILE
    if not path.is_file():
        raise FileNotFoundError(
            f"signal mapping table not found at {path}. "
            "Run `python scripts/fetch_benchmark.py` first."
        )

    instruments, derived_quality = _load_instruments(instruments_file)
    corrections_by_channel: dict[str, list[Correction]] = {}
    for correction in CORRECTIONS:
        corrections_by_channel.setdefault(correction.key, []).append(correction)

    signals: list[Signal] = []
    applied: list[Correction] = []

    for row in _read_rows(path):
        channel = _clean(row["Name"])
        if not channel:
            continue

        sensor_id = _clean(row["Sensor_ID"])
        for correction in corrections_by_channel.get(channel, ()):
            if correction.field_name == "sensor_id":
                if sensor_id != correction.was:
                    raise ValueError(
                        f"correction for {channel!r} expected sensor_id {correction.was!r} "
                        f"but the mapping table now says {sensor_id!r}. "
                        "The upstream file changed — review docs/benchmark-deviations.md."
                    )
                sensor_id = correction.now
                applied.append(correction)

        device_type = DeviceType(_clean(row["Type_of_Sensor_or_Actuator"]).lower())
        instrument = instruments.get(sensor_id)
        group = _classify_group(channel, device_type)
        quality, quality_reason = _resolve_quality(
            sensor_id, group, instrument, derived_quality, instruments
        )

        signals.append(
            Signal(
                index=int(row["ID"]),
                sensor_id=sensor_id,
                role=Role(_clean(row["Role"]).capitalize()),
                kind=SignalKind(_clean(row["Binary_or_Analogue"]).lower()),
                device_type=device_type,
                channel=channel,
                display_name=_DISPLAY_NAMES.get(channel, channel.replace("_", " ")),
                opcua_node_id=_clean(row["Sensor_OPCUA_Node_ID"]),
                unit=_nullable(row["Unit_Sensor"]),
                sim_variable=_nullable(row["Simulation_variable_name"]),
                sim_unit=_nullable(row["Unit_Simulation_variable"]),
                sim_init_parameter=_nullable(row["Simulation_initialization_parameters"]),
                group=group,
                asset_tag=_asset_tag(channel, sensor_id),
                quality=quality,
                quality_reason=quality_reason,
                instrument=instrument,
            )
        )

    missing = {c.key for c in CORRECTIONS} - {s.channel for s in signals}
    if missing:
        raise ValueError(
            f"corrections declared for channels that are not in the mapping table: {missing}"
        )

    return SignalDictionary(signals=tuple(signals), corrections_applied=tuple(applied))
